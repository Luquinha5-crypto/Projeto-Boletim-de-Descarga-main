"""Extrai do formulário ORIGINAL (xlsx) a geometria e o estilo de cada célula -> layout_original.json.

Rode só quando o formulário mudar:
    python tools/extrair_layout.py caminho/Descarga_Auto_Tanque_Pool.xlsx [layout_original.json]
(Se o original for .XLS, converta antes para .xlsx abrindo no Excel/LibreOffice e salvando como xlsx.)

O renderizador (gerar_pdf.py) só lê o JSON: assim o PDF segue o desenho do original sem
depender do Excel e sem reescrever coordenadas à mão.
"""
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ABA = "cálculo"
LINHAS, COLUNAS = 51, 18                      # área de impressão A1:R51
EMU_POR_PT = 12700
# Células do lado esquerdo que recebem DADOS (o original traz um exemplo preenchido nelas).
# Seus valores NÃO são copiados para o JSON: só o estilo. O lado direito espelha por fórmula (=B13...).
ENTRADAS = {"B12", "B13", "B14", "A17", "C17", "C18", "C20", "C21", "D21", "B35",
            "B25", "D25", "B27", "D27", "B29", "D29", "B31", "D31", "B33", "A38", "C38", "B45", "B47"}
ENTRADAS |= {f"{col}{lin}" for lin in (22, 23) for col in "ABCDEFGH" if (col, lin) != ("A", 22)}
CORRECOES_TEXTO = {                           # rótulos quebrados por espaços forçados no original
    "Conduti vidade": "Condutividade",
    "TemperaturaTanque": "Temperatura Tanque",
}


def pt_coluna(chars: float) -> float:
    """Largura de coluna do Excel (em caracteres) -> pontos."""
    return (chars * 7 + 5) * 0.75


def normalizar(texto):
    """Espaços em excesso no original são quebras de linha forçadas: 'Teor      Alcoólico' -> 2 linhas."""
    if not isinstance(texto, str) or texto.startswith("="):
        return texto
    t = re.sub(r"[ \t]{3,}", "\n", texto.strip())
    t = re.sub(r"[ \t]+", " ", t)
    return CORRECOES_TEXTO.get(t.replace("\n", " "), t)


def main(xlsx: Path, saida: Path):
    ws = openpyxl.load_workbook(xlsx)[ABA]

    larguras = [None] * COLUNAS
    for dim in ws.column_dimensions.values():
        for c in range(dim.min, dim.max + 1):
            if c <= COLUNAS:
                larguras[c - 1] = pt_coluna(dim.width)
    larguras = [w if w else pt_coluna(ws.sheet_format.defaultColWidth or 8.43) for w in larguras]
    alturas = [ws.row_dimensions[r].height or ws.sheet_format.defaultRowHeight for r in range(1, LINHAS + 1)]

    celulas = {}
    for r in range(1, LINHAS + 1):
        for c in range(1, COLUNAS + 1):
            x = ws.cell(r, c)
            lados = {k: getattr(x.border, n).style for k, n in
                     (("l", "left"), ("r", "right"), ("t", "top"), ("b", "bottom"))
                     if getattr(x.border, n).style}
            endereco = f"{get_column_letter(c)}{r}"
            valor = None if endereco in ENTRADAS else normalizar(x.value)
            if valor is not None and not isinstance(valor, str):
                valor = None          # números/datas soltos são dados de exemplo, nunca rótulos
            if not lados and valor is None and endereco not in ENTRADAS:
                continue                  # (células de entrada ficam sempre: precisam do estilo p/ escrever o dado)
            celulas[endereco] = {
                "v": valor,
                "fonte": [x.font.name, x.font.sz, bool(x.font.b), bool(x.font.i)],
                "alin": [x.alignment.horizontal, x.alignment.vertical, bool(x.alignment.wrap_text)],
                "borda": lados,
            }

    mescladas = [[m.min_row, m.min_col, m.max_row, m.max_col] for m in ws.merged_cells.ranges
                 if m.max_row <= LINHAS and m.max_col <= COLUNAS]

    # posição das imagens (logotipo) a partir do desenho da planilha
    imagens = []
    ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(xlsx) as z:
        for nome in z.namelist():
            if nome.startswith("xl/drawings/drawing") and nome.endswith(".xml"):
                raiz = ET.fromstring(z.read(nome))
                for anc in raiz.findall("xdr:twoCellAnchor", ns) + raiz.findall("xdr:oneCellAnchor", ns):
                    if anc.find("xdr:pic", ns) is None:
                        continue
                    de = anc.find("xdr:from", ns)
                    col, dcol = int(de.find("xdr:col", ns).text), int(de.find("xdr:colOff", ns).text)
                    lin, dlin = int(de.find("xdr:row", ns).text), int(de.find("xdr:rowOff", ns).text)
                    ext = anc.find("xdr:pic/xdr:spPr/a:xfrm/a:ext", ns)
                    imagens.append({
                        "x": sum(larguras[:col]) + dcol / EMU_POR_PT,
                        "y": sum(alturas[:lin]) + dlin / EMU_POR_PT,
                        "w": int(ext.get("cx")) / EMU_POR_PT, "h": int(ext.get("cy")) / EMU_POR_PT})

    # retângulos (as caixinhas DENTRO / FORA) a partir do desenho da planilha
    retangulos = []
    with zipfile.ZipFile(xlsx) as z:
        for nome in z.namelist():
            if nome.startswith("xl/drawings/drawing") and nome.endswith(".xml"):
                raiz = ET.fromstring(z.read(nome))
                for anc in raiz.findall("xdr:twoCellAnchor", ns) + raiz.findall("xdr:oneCellAnchor", ns):
                    forma = anc.find("xdr:sp", ns)
                    geo = None if forma is None else forma.find(".//a:prstGeom", ns)
                    if geo is None or geo.get("prst") != "rect":
                        continue
                    de = anc.find("xdr:from", ns)
                    col, dcol = int(de.find("xdr:col", ns).text), int(de.find("xdr:colOff", ns).text)
                    lin, dlin = int(de.find("xdr:row", ns).text), int(de.find("xdr:rowOff", ns).text)
                    ext = forma.find(".//a:xfrm/a:ext", ns)
                    retangulos.append({"x": sum(larguras[:col]) + dcol / EMU_POR_PT,
                                       "y": sum(alturas[:lin]) + dlin / EMU_POR_PT,
                                       "w": int(ext.get("cx")) / EMU_POR_PT, "h": int(ext.get("cy")) / EMU_POR_PT})

    dados = {"larguras": larguras, "alturas": alturas, "mescladas": mescladas, "celulas": celulas,
             "imagens": imagens, "retangulos": retangulos,
             "pagina": {"margem_lateral": ws.page_margins.left * 72, "margem_vertical": 5.0}}
    saida.write_text(json.dumps(dados, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(celulas)} células, {len(mescladas)} mescladas, {len(imagens)} imagens, {len(retangulos)} retângulos -> {saida}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) > 2 else Path("layout_original.json"))
