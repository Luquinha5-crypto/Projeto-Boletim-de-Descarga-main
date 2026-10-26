"""Gera a Autorização de Descarga em Excel (.xlsx) EDITÁVEL, igual ao formulário original.

Para que serve: o PDF sai pronto, mas o que a NF não trouxe (temperatura da amostra, pH, motorista...)
você pode digitar direto na planilha, sem JSON e sem --extra. A via da direita é feita de FÓRMULAS
(=B13, =A17...), como no original: o que você digita na via da esquerda aparece sozinho na da direita.

Reaproveita layout_original.json (o mesmo que desenha o PDF), então os dois ficam iguais.
"""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU

from calculos import Autorizacao
from gerar_pdf import LACRE_ENDERECOS, LAYOUT, LOGO
from volume import volume_20_exibido

PT_EMU = 12700                      # 1 ponto = 12700 EMU (unidade interna do Excel para desenhos)
FORMATOS = {                        # formato de número de cada célula de dado (o Excel mostra com a vírgula do seu Windows)
    "B25": "0.0", "D25": "0.0", "B27": "0.0000", "D27": "0.0000", "B29": "0.00", "D29": "0.00",
    "B31": "0.00", "D31": "General", "B33": "0.0000", "A38": "#,##0", "C38": "#,##0",
    "B12": "dd/mm/yy hh:mm",
}


def _largura_excel(pontos: float) -> float:
    """Inverso do que o PDF faz: largura em pontos -> largura em 'caracteres' do Excel."""
    return max(0.0, (pontos / 0.75 - 5) / 7)


def _formula_volume_20(a: Autorizacao) -> str:
    """Volume a 20 ºC como FÓRMULA (ambiente x fator), igual ao original: se você digitar o volume ambiente (A38)
    ou o fator (B33), ele recalcula sozinho. Sem esses dois, usa o valor da NF (ou a soma das NFs juntadas)."""
    reserva = a.volume_20_nf if a.volume_20_nf is not None else '""'
    return f"=IF(AND(ISNUMBER(A38),ISNUMBER(B33)),ROUND(A38*B33,0),{reserva})"


def dados_por_celula(a: Autorizacao) -> dict:
    """Endereço da via esquerda -> valor (número como número, texto como texto). Vazio = None."""
    carretas = (list(a.placas_carreta or []) + [None, None])[:2]
    v = {
        "B12": a.emitido_em, "B13": a.usina, "B14": a.produto, "A17": a.nota_fiscal, "C17": a.data_nf,
        "C18": a.transportador, "C20": a.placa_cavalo, "C21": carretas[0], "D21": carretas[1],
        "B25": a.temp_tanque, "D25": a.temp_amostra, "B27": a.dens_ambiente, "D27": a.dens_20,
        "B29": a.teor_alcoolico, "D29": a.acidez, "B31": a.ph, "D31": a.condutividade,
        "B33": a.fator_nf, "A38": a.volume_ambiente, "C38": _formula_volume_20(a),
        "B45": a.motorista, "B47": a.cpf,
    }
    for i, endereco in enumerate(LACRE_ENDERECOS):
        v[endereco] = str(a.lacres[i]) if i < len(a.lacres or []) else None   # texto: preserva zeros à esquerda
    return v


def gerar_xlsx(a: Autorizacao, caminho: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Autorização"
    ws.sheet_view.showGridLines = False

    for i, w in enumerate(LAYOUT["larguras"], start=1):
        ws.column_dimensions[get_column_letter(i)].width = _largura_excel(w)
    for i, h in enumerate(LAYOUT["alturas"], start=1):
        ws.row_dimensions[i].height = h

    dados = dados_por_celula(a)
    for endereco, cel in LAYOUT["celulas"].items():
        c = ws[endereco]
        nome, tamanho, negrito, italico = cel["fonte"]
        c.font = Font(name=nome, size=tamanho, bold=negrito, italic=italico)
        h, v, quebra = cel["alin"]
        c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=quebra)
        lados = {lado: Side(style=cel["borda"].get(ch)) for lado, ch in
                 (("left", "l"), ("right", "r"), ("top", "t"), ("bottom", "b"))}
        c.border = Border(**lados)

        valor = cel.get("v")                      # rótulo fixo ou fórmula do original (=B13, =A17...)
        if endereco in dados:                     # célula de dado da via esquerda
            valor = dados[endereco]
            if endereco in FORMATOS:
                c.number_format = FORMATOS[endereco]
        elif isinstance(valor, str) and valor.startswith("="):
            # via direita: espelha a esquerda. "=NOW()" vira "=B12" (mesma data/hora nas duas vias).
            origem = "B12" if valor.upper() == "=NOW()" else valor[1:].lstrip("+")
            if origem in dados:
                # SE(origem vazia; vazio; origem): sem isso a célula vazia apareceria como 0 (bug do original)
                valor = f'=IF({origem}="","",{origem})'
                if origem in FORMATOS:
                    c.number_format = FORMATOS[origem]             # mesmo formato do espelhado
        if valor is not None:
            c.value = valor

    for r1, c1, r2, c2 in LAYOUT["mescladas"]:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)

    if LOGO.exists():                             # logotipo opcional (assets/logo.jpeg, fora do Git)
        for img in LAYOUT["imagens"]:
            figura = Image(str(LOGO))
            figura.anchor = AbsoluteAnchor(
                pos=XDRPoint2D(int(img["x"] * PT_EMU), int(img["y"] * PT_EMU)),
                ext=XDRPositiveSize2D(int(img["w"] * PT_EMU), int(img["h"] * PT_EMU)))
            ws.add_image(figura)

    # impressão: A4 paisagem, tudo em UMA folha, centralizado
    ws.print_area = f"A1:{get_column_letter(len(LAYOUT['larguras']))}{len(LAYOUT['alturas'])}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    margem_h = LAYOUT["pagina"]["margem_lateral"] / 72
    margem_v = LAYOUT["pagina"]["margem_vertical"] / 72
    ws.page_margins.left = ws.page_margins.right = margem_h
    ws.page_margins.top = ws.page_margins.bottom = margem_v
    ws.page_margins.header = ws.page_margins.footer = 0

    wb.save(caminho)
