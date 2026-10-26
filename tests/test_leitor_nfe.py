"""Testes do leitor de NF-e com XMLs FICTÍCIOS (tests/fixtures/). Rodam em qualquer máquina e podem ir ao GitHub.

Rode na pasta do projeto:   python -m unittest discover -s tests -v
Os fixtures imitam o formato de 3 usinas reais (biodiesel, hidratado, anidro) com dados inventados.
Um teste FALHA de propósito (placas do hidratado): ele aponta um bug real do leitor.
"""
import dataclasses
import sys
import tempfile
import unittest
from pathlib import Path

RAIZ = Path(__file__).parent.parent
sys.path.insert(0, str(RAIZ))
from gerar import processar, processar_junto  # noqa: E402
from juntar import conflitos, divergencias, juntar, sugerir_grupos  # noqa: E402
from diagnostico import verificar  # noqa: E402
from leitor_nfe import carregar_cavalos, ler_nfe  # noqa: E402

FIX = Path(__file__).parent / "fixtures"


def ler(nome):
    return ler_nfe(str(FIX / f"ficticia_{nome}.xml"))


class LeitorNFe(unittest.TestCase):
    def test_fixtures_existem(self):  # evita o "passou sem rodar nada"
        self.assertEqual(len(list(FIX.glob("*.xml"))), 11)

    def test_campos_fixos(self):
        casos = [("biodiesel", "100001", "10/10/2026", "USINA BIO EXEMPLO S.A.", 50000),
                 ("hidratado", "100002", "13/10/2026", "USINA ALCOOL EXEMPLO S.A.", 40000),
                 ("anidro", "100003", "13/10/2026", "USINA ANIDRO EXEMPLO S.A.", 49900)]
        for nome, nf, data, usina, vol in casos:
            with self.subTest(nome):
                a, _ = ler(nome)
                self.assertEqual((a.nota_fiscal, a.data_nf, a.usina, a.volume_20_nf), (nf, data, usina, vol))

    def test_formato_abreviado(self):
        a, _ = ler("abreviada")
        self.assertEqual((a.temp_tanque, a.temp_amostra, a.fator_nf, a.teor_alcoolico), (22, 22, 0.9979, 93.1))
        self.assertEqual((a.dens_ambiente, a.dens_20), (0.8077, 0.8094))
        self.assertEqual((a.volume_ambiente, a.volume_20_nf), (42003, 41915))  # ponto = milhar
        self.assertEqual((a.ph, a.acidez, a.condutividade), (7.2, 6.38, 75.6))
        self.assertEqual((a.motorista, a.cpf), ("FULANO DE TAL", "000.000.000-00"))
        self.assertEqual(a.lacres, ["1000001", "1000002", "1000003", "1000004"])
        self.assertEqual((a.placa_cavalo, a.placas_carreta), ("AAA1B23", ["BBB4567"]))
        self.assertEqual(verificar(a, (FIX / "ficticia_abreviada.xml").read_text(encoding="utf-8")), [])

    def test_anidro_texto_livre(self):
        a, _ = ler("anidro")
        self.assertEqual((a.temp_tanque, a.fator_nf, a.teor_alcoolico), (21.9, 0.998, 99.6))
        self.assertEqual((a.volume_ambiente, a.volume_20_nf), (50000, 49900))
        self.assertEqual(a.lacres, ["0001001", "0001002", "0001003", "0001004"])
        self.assertEqual((a.placa_cavalo, a.placas_carreta), ("ABC1D23", ["ABC4E56"]))

    def test_lacre_de_amostra_nunca_entra(self):
        self.assertNotIn("200001", ler("biodiesel")[0].lacres)
        self.assertNotIn("E0009999", ler("hidratado")[0].lacres)
        self.assertNotIn("3000001", ler("anidro")[0].lacres)

    def test_campo_ausente_vira_aviso_e_nao_valor_chutado(self):
        a, avisos = ler("hidratado")     # esta NF não traz temperatura nem densidade
        self.assertIsNone(a.temp_tanque)
        self.assertTrue(any("temp_tanque" in v for v in avisos))

    def test_placas_do_hidratado(self):
        """BUG REAL: o texto traz 'CPF12345678909 CNH1234567890'. Veja o que o leitor devolve."""
        a, _ = ler("hidratado")
        self.assertEqual(a.placa_cavalo, "ABC1D23")
        self.assertEqual(a.placas_carreta, ["ABC4E56"])


class JuntarCasoBiodiesel(unittest.TestCase):
    """Par no formato do caso real (biodiesel): a NF A traz qualidade e lacres; a B traz só placa e volume."""

    def par(self):
        return ler("bio_a")[0], ler("bio_b")[0]

    def test_completa_e_soma_o_volume(self):
        a, b = self.par()
        avisos = []
        u = juntar([a, b], avisos)
        self.assertEqual((u.temp_tanque, u.dens_20, u.fator_nf, len(u.lacres)), (37, 0.878, 0.9871, 8))
        self.assertEqual((u.nota_fiscal, u.volume_20_nf), ("100010 / 100011", 56000))
        for campo in ("temp_tanque", "dens_20", "fator_nf", "lacres"):
            self.assertTrue(any(f"COMPLETADO {campo} veio só da NF 100010" in v and "na NF 100011" in v
                                for v in avisos), campo)

    def test_biodiesel_nao_gera_aviso_de_so_alcool(self):
        with tempfile.TemporaryDirectory() as tmp:
            import contextlib
            import io
            saida = io.StringIO()
            with contextlib.redirect_stdout(saida):
                processar_junto([FIX / "ficticia_bio_a.xml", FIX / "ficticia_bio_b.xml"], Path(tmp), None)
            self.assertNotIn("para álcool", saida.getvalue())

    def test_nao_sugere_se_os_lacres_forem_diferentes(self):
        a, b = self.par()
        b = dataclasses.replace(b, lacres=["999999"])        # outra carga no mesmo caminhão e dia
        self.assertEqual(sugerir_grupos([a, b]), [])

    def test_gera_um_pdf_e_um_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            codigo = processar_junto([FIX / "ficticia_bio_a.xml", FIX / "ficticia_bio_b.xml"], Path(tmp), None)
            self.assertEqual(codigo, 0)
            self.assertEqual(sorted(p.name for p in Path(tmp).iterdir()),
                             ["boletim_NF100010_100011.pdf", "boletim_NF100010_100011.xlsx"])


class JuntarLacres(unittest.TestCase):
    """Regra do usuário: números iguais aparecem UMA vez (união); números diferentes = conflito."""

    def par(self):
        a, b = ler("bio_a")[0], ler("bio_b")[0]
        return a, dataclasses.replace(b, placa_cavalo=a.placa_cavalo)

    def duas_com_lacres(self, lacres_b):
        a, b = self.par()
        return a, dataclasses.replace(b, lacres=lacres_b)

    def test_lacres_iguais_aparecem_uma_vez(self):
        a, b = self.duas_com_lacres(list(ler("bio_a")[0].lacres))
        self.assertEqual(juntar([a, b]).lacres, a.lacres)
        self.assertEqual(len(juntar([a, b]).lacres), 8)

    def test_uma_nf_com_parte_dos_lacres_une_sem_repetir(self):
        a, b = self.duas_com_lacres(ler("bio_a")[0].lacres[:3])
        avisos = []
        u = juntar([a, b], avisos)
        self.assertEqual(u.lacres, a.lacres)               # os 8, sem repetir os 3 que estavam nas duas
        self.assertTrue(any("COMPLETADO lacres veio só da NF 100010" in v for v in avisos))

    def test_lacres_totalmente_diferentes_sao_conflito(self):
        a, b = self.duas_com_lacres(["111111", "222222"])
        self.assertEqual(list(conflitos([a, b])), ["lacres"])
        with self.assertRaises(ValueError):
            juntar([a, b])

    def test_lacres_parcialmente_diferentes_sao_conflito(self):
        a, b = self.duas_com_lacres(ler("bio_a")[0].lacres[:2] + ["999999"])   # 2 iguais + 1 que a outra não tem
        self.assertIn("lacres", conflitos([a, b]))

    def test_sugestao_aceita_lacres_em_parte_e_recusa_diferentes(self):
        a, b = self.duas_com_lacres(ler("bio_a")[0].lacres[:3])
        self.assertEqual(len(sugerir_grupos([a, b])), 1)
        a, b = self.duas_com_lacres(["111111"])
        self.assertEqual(sugerir_grupos([a, b]), [])


class LacresLaterais(unittest.TestCase):
    def test_lacres_LA_nao_entram_no_boletim(self):
        a, _ = ler("biodiesel")   # o texto tem "Lacre:...-100000*LA:100010-100009-100008"
        self.assertEqual(len(a.lacres), 8)
        for lateral in ("100010", "100009", "100008"):
            self.assertNotIn(lateral, a.lacres)


class ListaDeCavalos(unittest.TestCase):
    def arquivo(self, texto):
        pasta = tempfile.TemporaryDirectory()
        self.addCleanup(pasta.cleanup)
        caminho = Path(pasta.name) / "cavalos.txt"
        caminho.write_text(texto, encoding="utf-8")
        return caminho

    def test_carregar_ignora_comentario_hifen_e_minuscula(self):
        c = carregar_cavalos(self.arquivo("# frota\nbbb-4567  # carreta?\n\n  AAA1B23\n"))
        self.assertEqual(c, {"BBB4567", "AAA1B23"})

    def test_lista_define_o_cavalo_mesmo_fora_da_ordem(self):
        # na NF a ordem é AAA1B23, BBB4567; a lista diz que o cavalo é o BBB4567
        a, avisos = ler_nfe(str(FIX / "ficticia_abreviada.xml"), {"BBB4567"})
        self.assertEqual((a.placa_cavalo, a.placas_carreta), ("BBB4567", ["AAA1B23"]))
        self.assertFalse(any("Placas" in v for v in avisos))

    def test_placa_fora_da_lista_cai_no_aviso(self):
        a, avisos = ler_nfe(str(FIX / "ficticia_abreviada.xml"), {"ZZZ9Z99"})
        self.assertEqual(a.placa_cavalo, "AAA1B23")
        self.assertTrue(any("nenhuma das placas" in v for v in avisos))


class JuntarNFs(unittest.TestCase):
    def duas(self):
        return ler("fracao_a")[0], ler("fracao_b")[0]

    def test_junta_soma_volume_e_mostra_as_duas_nfs(self):
        a, b = self.duas()
        u = juntar([a, b])
        self.assertEqual(u.nota_fiscal, "100006 / 100007")
        self.assertEqual(u.volume_20_nf, 41915)
        self.assertEqual(u.volume_ambiente, 42003)     # capacidade: NÃO soma
        self.assertEqual(u.data_nf, "14/10/2026")

    def test_aceita_mais_de_duas(self):
        a, b = self.duas()
        c = dataclasses.replace(a, nota_fiscal="100008", volume_20_nf=100)
        self.assertEqual(juntar([a, b, c]).volume_20_nf, 42015)

    def test_motorista_diferente_nao_junta(self):
        a, b = self.duas()
        b = dataclasses.replace(b, motorista="OUTRO MOTORISTA")
        self.assertEqual(list(divergencias([a, b])), ["motorista"])
        with self.assertRaises(ValueError) as e:
            juntar([a, b])
        self.assertIn("motorista", str(e.exception))

    def test_conflito_real_com_um_vazio_so_conta_o_que_diverge(self):
        a, b = self.duas()
        b = dataclasses.replace(b, motorista="OUTRO", lacres=[])        # lacres vazio NÃO é conflito
        self.assertEqual(list(conflitos([a, b])), ["motorista"])

    def test_nf_repetida_nao_junta(self):
        a, _ = self.duas()
        with self.assertRaises(ValueError):
            juntar([a, a])

    def test_sugere_grupo_pelos_mesmos_lacres_e_placas(self):
        a, b = self.duas()
        outra = ler("anidro")[0]
        grupos = sugerir_grupos([a, b, outra])
        self.assertEqual(len(grupos), 1)
        self.assertEqual({x.nota_fiscal for x in grupos[0]}, {"100006", "100007"})

    def test_conferencia_fecha_na_soma(self):
        """Cada NF sozinha tem só parte do volume (confere mal); a soma bate com capacidade x fator."""
        from calculos import conferir
        u = juntar(list(self.duas()))
        vol = [v for v in conferir(u) if v["item"].startswith("Volume")][0]
        self.assertTrue(vol["ok"])

    def test_gera_um_pdf_so(self):
        with tempfile.TemporaryDirectory() as tmp:
            codigo = processar_junto([FIX / "ficticia_fracao_a.xml", FIX / "ficticia_fracao_b.xml"],
                                     Path(tmp), None)
            self.assertEqual(codigo, 0)
            self.assertEqual([p.name for p in Path(tmp).glob("*.pdf")], ["boletim_NF100006_100007.pdf"])


class PlanilhaEditavel(unittest.TestCase):
    """O .xlsx que sai junto do PDF: dados nas células certas, via direita espelhando, sem '0' falso."""

    def setUp(self):
        from openpyxl import load_workbook
        from gerar_xlsx import gerar_xlsx
        pasta = tempfile.TemporaryDirectory()
        self.addCleanup(pasta.cleanup)
        self.caminho = Path(pasta.name) / "b.xlsx"
        gerar_xlsx(ler("abreviada")[0], str(self.caminho))
        self.ws = load_workbook(self.caminho).active

    def test_dados_nas_celulas_certas_com_tipos(self):
        ws = self.ws
        self.assertEqual((ws["B13"].value, ws["A17"].value, ws["C20"].value), ("USINA ABREVIADA LTDA", "100005", "AAA1B23"))
        self.assertEqual((ws["B22"].value, ws["E22"].value), ("1000001", "1000004"))
        self.assertIsNone(ws["F22"].value)                                                         # sem lacre = vazio

    def test_via_direita_espelha_com_formula_e_sem_zero_falso(self):
        self.assertEqual(self.ws["L13"].value, '=IF(B13="","",B13)')
        self.assertEqual(self.ws["L22"].value, '=IF(B22="","",B22)')
        self.assertEqual(self.ws["L12"].value, '=IF(B12="","",B12)')   # data/hora: espelha, não NOW()

    def test_layout_completo(self):
        self.assertEqual(len(self.ws.merged_cells.ranges), 152)
        self.assertEqual(self.ws.page_setup.orientation, "landscape")

    def test_gerar_cria_pdf_e_xlsx_ou_so_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            processar(FIX / "ficticia_abreviada.xml", Path(tmp), None)
            self.assertEqual(sorted(p.suffix for p in Path(tmp).iterdir()), [".pdf", ".xlsx"])
        with tempfile.TemporaryDirectory() as tmp:
            processar(FIX / "ficticia_abreviada.xml", Path(tmp), None, com_xlsx=False)
            self.assertEqual([p.suffix for p in Path(tmp).iterdir()], [".pdf"])


class JuntarComDadosFaltando(unittest.TestCase):
    """Opção A: campo vazio numa NF e preenchido na outra é COMPLETADO (com aviso). Conflito real continua bloqueando."""

    def par(self):
        return ler("completa")[0], ler("esparsa")[0]

    def test_o_caso_real_esparsa_nao_trazia_qualidade_nem_lacres(self):
        completa, esparsa = self.par()
        self.assertEqual((esparsa.temp_tanque, esparsa.fator_nf, esparsa.lacres), (None, None, []))
        self.assertEqual(completa.temp_tanque, 22)

    def test_completa_os_campos_vazios_e_avisa(self):
        for ordem in ("normal", "invertida"):          # não pode importar qual NF vem primeiro
            with self.subTest(ordem):
                a, b = self.par()
                avisos = []
                u = juntar([a, b] if ordem == "normal" else [b, a], avisos)
                self.assertEqual((u.temp_tanque, u.fator_nf, u.dens_20, u.volume_ambiente), (22, 0.9979, 0.8094, 42003))
                self.assertEqual(u.lacres, ["1000001", "1000002", "1000003", "1000004"])
                self.assertEqual(u.volume_20_nf, 41915)
                self.assertTrue(any("lacres veio só da NF 100008" in v for v in avisos))

    def test_placa_escrita_diferente_nao_e_conflito(self):
        a, b = self.par()          # a esparsa traz "AAA-1B23" (com hífen); a outra, "AAA1B23"
        self.assertEqual(b.placa_cavalo, "AAA1B23")   # o leitor já tira o hífen? se não, a normalização cobre
        self.assertEqual(conflitos([a, b]), {})

    def test_sem_identificador_forte_nao_completa(self):
        a, b = self.par()
        b = dataclasses.replace(b, placa_cavalo=None, placas_carreta=[])   # nada liga a NF ao caminhão
        with self.assertRaises(ValueError) as e:
            juntar([a, b], [])
        self.assertIn("identificador", str(e.exception))

    def test_conflito_real_continua_bloqueando(self):
        a, b = self.par()
        b = dataclasses.replace(b, temp_tanque=30.0)      # a outra NF diz 22
        with self.assertRaises(ValueError) as e:
            juntar([a, b])
        self.assertIn("temp_tanque", str(e.exception))

    def test_sugere_grupo_mesmo_sem_lacre_em_uma_das_nfs(self):
        a, b = self.par()
        self.assertEqual(len(sugerir_grupos([a, b])), 1)

    def test_conferencia_fecha_na_soma_completada(self):
        from calculos import conferir
        u = juntar(list(self.par()), [])
        self.assertNotIn(False, [v["ok"] for v in conferir(u)])


class Diagnostico(unittest.TestCase):
    """O alerta de leitura: cala quando o leitor entendeu, avisa quando o texto cita o dado e ele não foi lido."""

    @staticmethod
    def suspeitas(nome):
        from xml.etree import ElementTree as ET
        a, _ = ler(nome)
        texto = ET.parse(FIX / f"ficticia_{nome}.xml").getroot().findtext(
            ".//{http://www.portalfiscal.inf.br/nfe}infCpl", "")
        return verificar(a, texto)

    def test_formatos_conhecidos_nao_geram_alarme_falso(self):
        for nome in ("anidro", "biodiesel", "hidratado"):
            with self.subTest(nome):
                self.assertEqual(self.suspeitas(nome), [])

    def test_formato_novo_gera_suspeitas(self):
        campos = " ".join(self.suspeitas("formato_novo"))
        for esperado in ("temp_tanque", "dens_20", "fator_nf", "teor_alcoolico"):
            self.assertIn(esperado, campos)


class DoXmlAoPdf(unittest.TestCase):
    def test_gera_pdf_para_cada_xml(self):
        with tempfile.TemporaryDirectory() as tmp:
            for xml in sorted(FIX.glob("*.xml")):
                with self.subTest(xml.name):
                    processar(xml, Path(tmp), None)
            pdfs = list(Path(tmp).glob("*.pdf"))
            self.assertEqual(len(pdfs), 11)
            self.assertTrue(all(p.stat().st_size > 5000 for p in pdfs))


if __name__ == "__main__":
    unittest.main()
